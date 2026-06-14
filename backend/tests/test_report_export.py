"""Экспорт дневного отчёта в CSV/XLSX (MVP-2)."""

import csv
import io
from datetime import date

import pytest

from app.adapters.telegram import MockTelegramClient
from app.core.enums import DayPart, PaymentMethod
from app.services import driver_actions
from app.services import orders as orders_service
from app.services.dispatch import send_batch
from app.services.errors import ValidationError
from app.services.planning import create_batch
from app.services.report_export import export_daily_report
from tests.conftest import order_payload


async def _completed_order(session, fx):
    order = await orders_service.create_order(session, order_payload(fx), actor_id=1)
    await orders_service.assign_driver(session, order, fx["driver"].id, actor_id=1)
    batch, _ = await create_batch(
        session, date.today(), DayPart.first_half, fx["district"].id, fx["driver"].id, [order.id]
    )
    await send_batch(session, batch, MockTelegramClient())
    driver = fx["driver"]
    await driver_actions.batch_accept(session, batch, driver)
    await driver_actions.batch_start(session, batch, driver)
    await driver_actions.order_complete(session, driver, order, PaymentMethod.cash, None)
    return order


async def test_export_csv(session, fixtures):
    await _completed_order(session, fixtures)
    content, media_type, filename = await export_daily_report(session, date.today(), "csv")

    assert "text/csv" in media_type
    assert filename.endswith(".csv")
    text = content.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[0][0] == "№"  # заголовок
    assert len(rows) == 2  # заголовок + 1 заказ
    assert rows[1][1] == "Иван Петров"  # клиент из снапшота
    assert "Выполнен" in rows[1]


async def test_export_xlsx_three_sheets(session, fixtures):
    await _completed_order(session, fixtures)
    content, media_type, filename = await export_daily_report(session, date.today(), "xlsx")

    assert "spreadsheetml" in media_type
    assert filename.endswith(".xlsx")

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Сводка", "Касса по водителям", "Заказы"]
    # касса содержит водителя с суммой
    cash_sheet = wb["Касса по водителям"]
    data = list(cash_sheet.iter_rows(values_only=True))
    assert data[0][0] == "Водитель"
    assert any(row[0] == "Сергей" for row in data[1:])
    # лист заказов содержит одну строку данных
    orders_sheet = wb["Заказы"]
    assert orders_sheet.max_row == 2


async def test_export_empty_day_has_headers_only(session, fixtures):
    content, _, _ = await export_daily_report(session, date(2020, 1, 1), "csv")
    rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    assert len(rows) == 1  # только заголовок
    assert rows[0][0] == "№"


async def test_export_invalid_format(session, fixtures):
    with pytest.raises(ValidationError, match="csv или xlsx"):
        await export_daily_report(session, date.today(), "pdf")
