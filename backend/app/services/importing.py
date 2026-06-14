"""Импорт заказов из CSV/XLSX (SPEC.md, раздел 3 MVP-2; эндпоинт раздела 8).

Один файл = много строк, каждая строка = заказ. Для каждой строки клиент и адрес
ищутся (find-or-create) по телефону и тексту адреса, затем создаётся заказ со
снапшотом (source_type=import). Ошибочные строки не валят импорт: они
откатываются по savepoint и попадают в отчёт, остальные сохраняются.

Режим dry_run — предпросмотр: всё парсится и валидируется, но в конце откат.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.enums import GeocodeStatus, PaymentMethodPlan, SourceType, TimeWindowType
from app.core.phones import normalize_phone as _normalize_phone
from app.models import Address, Client, District
from app.services.geocoding import normalize_address
from app.services.orders import create_order

# Синонимы заголовков (нормализованные: нижний регистр, схлопнутые пробелы)
HEADER_ALIASES: dict[str, set[str]] = {
    "name": {"имя", "клиент", "фио", "name", "client"},
    "phone": {"телефон", "тел", "phone"},
    "address": {"адрес", "address", "raw_address"},
    "district": {"район", "district"},
    "latitude": {"широта", "lat", "latitude"},
    "longitude": {"долгота", "lon", "lng", "longitude"},
    "entrance": {"подъезд", "entrance"},
    "floor": {"этаж", "floor"},
    "date": {"дата", "date", "delivery_date"},
    "day_part": {"часть дня", "часть_дня", "окно", "day_part", "time_window"},
    "bottles_pc": {"бутыли пк", "бутыли_пк", "пк", "поликарбонат", "bottles_pc"},
    "bottles_pet": {"бутыли пэт", "бутыли_пэт", "пэт", "bottles_pet"},
    "pumps": {"помпы", "помпа", "pumps"},
    "amount": {"сумма", "amount", "total", "total_amount"},
    "payment": {"оплата", "способ оплаты", "payment", "payment_method_plan"},
    "comment": {"комментарий", "коммент", "примечание", "comment"},
}

_DAY_PART = {
    "": TimeWindowType.any,
    "любое": TimeWindowType.any,
    "любая": TimeWindowType.any,
    "any": TimeWindowType.any,
    "первая": TimeWindowType.first_half,
    "первая половина": TimeWindowType.first_half,
    "до обеда": TimeWindowType.first_half,
    "first_half": TimeWindowType.first_half,
    "вторая": TimeWindowType.second_half,
    "вторая половина": TimeWindowType.second_half,
    "после обеда": TimeWindowType.second_half,
    "second_half": TimeWindowType.second_half,
    "точное": TimeWindowType.exact,
    "exact": TimeWindowType.exact,
}

_PAYMENT = {
    "": PaymentMethodPlan.unknown,
    "не известно": PaymentMethodPlan.unknown,
    "уточнить": PaymentMethodPlan.unknown,
    "unknown": PaymentMethodPlan.unknown,
    "наличные": PaymentMethodPlan.cash,
    "нал": PaymentMethodPlan.cash,
    "cash": PaymentMethodPlan.cash,
    "карта": PaymentMethodPlan.cashless,
    "перевод": PaymentMethodPlan.cashless,
    "карта/перевод": PaymentMethodPlan.cashless,
    "безнал": PaymentMethodPlan.cashless,
    "cashless": PaymentMethodPlan.cashless,
}


class ImportError_(Exception):
    """Ошибка структуры файла (нечитаемый формат, нет обязательных колонок)."""


@dataclass
class ImportReport:
    total_rows: int = 0
    imported: int = 0
    clients_created: int = 0
    addresses_created: int = 0
    errors: list[dict] = field(default_factory=list)
    dry_run: bool = False
    order_ids: list[int] = field(default_factory=list)


def _norm_header(h: str) -> str:
    return re.sub(r"\s+", " ", (h or "").strip().lower())


def _map_headers(raw_headers: list[str]) -> dict[str, int]:
    """Канонический ключ -> индекс колонки."""
    mapping: dict[str, int] = {}
    for idx, raw in enumerate(raw_headers):
        key = _norm_header(raw)
        for canon, aliases in HEADER_ALIASES.items():
            if key in aliases and canon not in mapping:
                mapping[canon] = idx
    return mapping


def parse_file(filename: str, content: bytes) -> list[dict]:
    """Возвращает список строк-словарей {канонический ключ: строковое значение}."""
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        rows = _read_xlsx(content)
    elif name.endswith(".csv") or name.endswith(".txt"):
        rows = _read_csv(content)
    else:
        raise ImportError_("Поддерживаются только файлы .csv и .xlsx")
    if not rows:
        raise ImportError_("Файл пустой")

    headers, *data_rows = rows
    mapping = _map_headers(headers)
    for required in ("name", "phone", "address", "date"):
        if required not in mapping:
            raise ImportError_(
                f"Не найдена обязательная колонка: {required} "
                f"(допустимые заголовки: {', '.join(sorted(HEADER_ALIASES[required]))})"
            )

    result = []
    for raw in data_rows:
        if not any((c or "").strip() for c in raw):
            continue  # пустая строка
        result.append(
            {canon: (raw[idx] if idx < len(raw) else "") for canon, idx in mapping.items()}
        )
    return result


def _read_csv(content: bytes) -> list[list[str]]:
    text = content.decode("utf-8-sig", errors="replace")
    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), dialect)
    return [[(c or "").strip() for c in row] for row in reader]


def _read_xlsx(content: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        rows.append([_cell_to_str(c) for c in row])
    wb.close()
    return rows


def _cell_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


# --- Парсинг значений строки (бросают ValueError с понятным сообщением) ---


def _parse_date(raw: str) -> date:
    raw = raw.strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y", "%d/%m/%Y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"не распознана дата '{raw}' (ожидается ГГГГ-ММ-ДД или ДД.ММ.ГГГГ)")


def _parse_int(raw: str, field_name: str) -> int:
    raw = (raw or "").strip()
    if raw == "":
        return 0
    try:
        return int(float(raw.replace(",", ".")))
    except ValueError:
        raise ValueError(f"{field_name}: не число '{raw}'") from None


def _parse_decimal(raw: str, field_name: str) -> Decimal:
    raw = (raw or "").strip().replace(" ", "").replace(",", ".")
    if raw == "":
        return Decimal("0.00")
    try:
        return Decimal(raw)
    except InvalidOperation:
        raise ValueError(f"{field_name}: не сумма '{raw}'") from None


def _parse_coord(raw: str, field_name: str) -> Decimal | None:
    raw = (raw or "").strip().replace(",", ".")
    if raw == "":
        return None
    try:
        return Decimal(raw)
    except InvalidOperation:
        raise ValueError(f"{field_name}: не координата '{raw}'") from None


async def _load_phone_index(session: AsyncSession) -> dict[str, Client]:
    """Индекс существующих клиентов по нормализованному телефону.

    Телефоны в базе хранятся с разделителями, поэтому матч по LIKE ненадёжен —
    сравниваем нормализованные значения в памяти (масштаб справочника небольшой).
    """
    clients = (await session.execute(select(Client))).scalars().all()
    index: dict[str, Client] = {}
    for c in clients:
        index.setdefault(_normalize_phone(c.phone_primary), c)
    return index


async def _find_or_create_client(
    session: AsyncSession, name: str, phone: str, phone_index: dict[str, Client]
) -> tuple[Client, bool]:
    norm = _normalize_phone(phone)
    if norm and norm in phone_index:
        return phone_index[norm], False
    client = Client(name=name, phone_primary=phone)
    session.add(client)
    await session.flush()
    if norm:
        phone_index[norm] = client
    return client, True


async def _find_or_create_address(
    session: AsyncSession, client: Client, row: dict, district_id: int | None
) -> tuple[Address, bool]:
    key = normalize_address(row["address"])
    existing = (
        (await session.execute(select(Address).where(Address.client_id == client.id)))
        .scalars()
        .all()
    )
    for a in existing:
        if normalize_address(a.raw_address) == key:
            return a, False
    lat = _parse_coord(row.get("latitude", ""), "широта")
    lon = _parse_coord(row.get("longitude", ""), "долгота")
    has_coords = lat is not None and lon is not None
    address = Address(
        client_id=client.id,
        raw_address=row["address"].strip(),
        district_id=district_id,
        latitude=lat,
        longitude=lon,
        geocode_status=GeocodeStatus.manual if has_coords else GeocodeStatus.none,
        entrance=row.get("entrance", "").strip() or None,
        floor=row.get("floor", "").strip() or None,
    )
    session.add(address)
    await session.flush()
    return address, True


async def _district_id_by_name(session: AsyncSession, name: str) -> int | None:
    name = (name or "").strip()
    if not name:
        return None
    district = (
        await session.execute(select(District).where(District.name.ilike(name)))
    ).scalar_one_or_none()
    if district is None:
        raise ValueError(f"район '{name}' не найден в справочнике")
    return district.id


async def _process_row(
    session: AsyncSession, row: dict, actor_id: int | None, phone_index: dict[str, Client]
) -> dict:
    """Создаёт клиента/адрес/заказ для строки. Бросает ValueError при проблемах."""
    name = (row.get("name") or "").strip()
    phone = (row.get("phone") or "").strip()
    address_text = (row.get("address") or "").strip()
    if not name:
        raise ValueError("пустое имя клиента")
    if not phone:
        raise ValueError("пустой телефон")
    if not address_text:
        raise ValueError("пустой адрес")

    delivery_date = _parse_date(row.get("date", ""))
    district_id = await _district_id_by_name(session, row.get("district", ""))

    day_part_raw = _norm_header(row.get("day_part", ""))
    if day_part_raw not in _DAY_PART:
        raise ValueError(f"часть дня: неизвестно '{day_part_raw}'")
    payment_raw = _norm_header(row.get("payment", ""))
    if payment_raw not in _PAYMENT:
        raise ValueError(f"оплата: неизвестно '{payment_raw}'")

    client, client_created = await _find_or_create_client(session, name, phone, phone_index)
    address, address_created = await _find_or_create_address(session, client, row, district_id)

    order = await create_order(
        session,
        {
            "client_id": client.id,
            "address_id": address.id,
            "delivery_date": delivery_date,
            "time_window_type": _DAY_PART[day_part_raw],
            "bottles_pc_qty": _parse_int(row.get("bottles_pc", ""), "бутыли ПК"),
            "bottles_pet_qty": _parse_int(row.get("bottles_pet", ""), "бутыли ПЭТ"),
            "pumps_qty": _parse_int(row.get("pumps", ""), "помпы"),
            "total_amount": _parse_decimal(row.get("amount", ""), "сумма"),
            "payment_method_plan": _PAYMENT[payment_raw],
            "comment": (row.get("comment") or "").strip() or None,
            "source_type": SourceType.import_,
        },
        actor_id=actor_id,
    )
    return {
        "order_id": order.id,
        "client_created": client_created,
        "address_created": address_created,
    }


async def import_orders(
    session: AsyncSession,
    filename: str,
    content: bytes,
    actor_id: int | None,
    dry_run: bool = False,
) -> ImportReport:
    rows = parse_file(filename, content)
    report = ImportReport(total_rows=len(rows), dry_run=dry_run)
    phone_index = await _load_phone_index(session)

    for i, row in enumerate(rows, start=2):  # строка 1 — заголовок
        try:
            async with session.begin_nested():  # savepoint: ошибочная строка откатывается
                res = await _process_row(session, row, actor_id, phone_index)
            report.imported += 1
            report.order_ids.append(res["order_id"])
            report.clients_created += int(res["client_created"])
            report.addresses_created += int(res["address_created"])
        except Exception as e:  # noqa: BLE001 — собираем все ошибки строк в отчёт
            report.errors.append({"row": i, "error": str(e)})

    if dry_run:
        await session.rollback()
        report.order_ids = []  # в предпросмотре id не сохраняются
    else:
        await session.commit()
    return report
