"""Экспорт дневного отчёта в CSV/XLSX (SPEC.md, раздел 8, MVP-2).

CSV — детальная таблица заказов за день (один заказ = одна строка).
XLSX — три листа: «Сводка», «Касса по водителям», «Заказы».
"""

import csv
import io
from datetime import date

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Driver, Order
from app.services.reporting import daily_report

CSV_MEDIA = "text/csv; charset=utf-8"
XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

STATUS_RU = {
    "draft": "Черновик",
    "new": "Новый",
    "needs_review": "Требует проверки",
    "planned": "Запланирован",
    "assigned": "Назначен",
    "sent_to_driver": "Отправлен водителю",
    "accepted_by_driver": "Принят водителем",
    "in_progress": "В работе",
    "completed": "Выполнен",
    "failed": "Ошибка",
    "refused": "Отказ",
    "postponed": "Перенос",
    "cancelled": "Отменён",
}
PART_RU = {"any": "Любое", "first_half": "Первая половина", "second_half": "Вторая половина",
           "exact": "Точное окно"}
PAY_STATUS_RU = {"unpaid": "Не оплачен", "paid": "Оплачен", "partial": "Частично"}
PAY_METHOD_RU = {"cash": "Наличные", "cashless": "Карта/перевод", "unknown": "Не известно"}

ORDER_HEADERS = [
    "№", "Клиент", "Телефон", "Адрес", "Район", "Часть дня",
    "Бутыли ПК", "Бутыли ПЭТ", "Помпы", "Сумма", "Оплачено",
    "Статус оплаты", "Статус", "Водитель", "Создан", "Выполнен",
]


async def _order_rows(session: AsyncSession, report_date: date) -> list[list]:
    orders = (
        (
            await session.execute(
                select(Order).where(Order.delivery_date == report_date).order_by(Order.id)
            )
        )
        .scalars()
        .all()
    )
    driver_names = {
        d.id: d.name for d in (await session.execute(select(Driver))).scalars().all()
    }
    # район берём из снапшота заказа? district_id — снапшот; имя нужно подтянуть
    from app.models import District

    district_names = {
        d.id: d.name for d in (await session.execute(select(District))).scalars().all()
    }
    rows = []
    for o in orders:
        rows.append(
            [
                o.id,
                o.client_name,
                o.phone,
                o.address_text,
                district_names.get(o.district_id, ""),
                PART_RU.get(o.time_window_type.value, o.time_window_type.value),
                o.bottles_pc_qty,
                o.bottles_pet_qty,
                o.pumps_qty,
                str(o.total_amount),
                str(o.paid_amount),
                PAY_STATUS_RU.get(o.payment_status.value, o.payment_status.value),
                STATUS_RU.get(o.status.value, o.status.value),
                driver_names.get(o.assigned_driver_id, ""),
                o.created_at.strftime("%Y-%m-%d %H:%M") if o.created_at else "",
                o.completed_at.strftime("%Y-%m-%d %H:%M") if o.completed_at else "",
            ]
        )
    return rows


def _to_csv(headers: list[str], rows: list[list]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    writer.writerows(rows)
    # BOM — чтобы кириллица корректно открывалась в Excel
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _to_xlsx(report: dict, order_headers: list[str], order_rows: list[list]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    bold = Font(bold=True)

    # Лист 1 — Сводка
    ws = wb.active
    ws.title = "Сводка"
    ws.append(["Дневной отчёт", report["date"]])
    ws["A1"].font = bold
    ws.append([])
    ws.append(["Заказы по статусам"])
    ws["A3"].font = bold
    for status, n in report["orders_by_status"].items():
        ws.append([STATUS_RU.get(status, status), n])
    ws.append(["Всего заказов", report["orders_total"]])
    ws.append([])
    b = report["bottles_delivered"]
    ws.append(["Бутыли (выполнено): ПК / ПЭТ / помпы", b["pc"], b["pet"], b["pumps"]])
    ws.append([])
    ws.append(["Деньги по способам оплаты"])
    ws[f"A{ws.max_row}"].font = bold
    for method, amount in report["money_by_method"].items():
        ws.append([PAY_METHOD_RU.get(method, method), str(amount)])

    # Лист 2 — Касса по водителям
    ws2 = wb.create_sheet("Касса по водителям")
    headers2 = ["Водитель", "Выполнено заказов", "Бутылей", "Наличные", "Карта/перевод"]
    ws2.append(headers2)
    for c in ws2[1]:
        c.font = bold
    for d in report["drivers"]:
        ws2.append([
            d["driver_name"], d["completed_orders"], d["bottles"],
            str(d["cash"]), str(d["cashless"]),
        ])

    # Лист 3 — Заказы
    ws3 = wb.create_sheet("Заказы")
    ws3.append(order_headers)
    for c in ws3[1]:
        c.font = bold
    for row in order_rows:
        ws3.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def export_daily_report(
    session: AsyncSession, report_date: date, fmt: str
) -> tuple[bytes, str, str]:
    """Возвращает (содержимое, media_type, имя файла)."""
    fmt = (fmt or "").lower()
    rows = await _order_rows(session, report_date)
    if fmt == "csv":
        return _to_csv(ORDER_HEADERS, rows), CSV_MEDIA, f"report_{report_date}.csv"
    if fmt == "xlsx":
        report = await daily_report(session, report_date)
        return _to_xlsx(report, ORDER_HEADERS, rows), XLSX_MEDIA, f"report_{report_date}.xlsx"
    from app.services.errors import ValidationError

    raise ValidationError("format должен быть csv или xlsx")
